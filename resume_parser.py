import os 
import time
from pathlib import Path
from dotenv import load_dotenv
from groq import Groq

env_path = Path.home() / "Desktop" / "notepad.env"
load_dotenv(env_path)
my_api_key = os.getenv("GROQ_API_KEY")

if not my_api_key:
    raise ValueError("where is api")

client = Groq(api_key=my_api_key)

model = "openai/gpt-oss-120b"

job_description = """
Description
The AWS Hardware Engineering team creates server designs for Amazon’s innovative cloud services. Our designs are industry-leading in frugality and operational excellence, and are critical to the success of the AWS business and millions of customers who use AWS today. Our engineers solve challenging technology problems, and build architecturally sound, high-quality components to enable AWS to realize critical business strategies.
The ideal candidate for this role will be a curious, innovative self-starter. You will be an experienced Firmware/Software Development Engineer specialized in programmable logic development with excellent hardware description language skills, knowledge in digital hardware and embedded firmware, and eager to expand your knowledge and skills. You will use broad knowledge of the system in your projects to find the best solutions to difficult problems. You will work with engineers across the company as well as external companies and lead development efforts to and beyond delivery. You will collaborate with internal and external development engineers (software, hardware, validation, software services). AWS Engineers are shaping the way people use computers and designing the future of cloud computing technology – come help us make history!

AWS Infrastructure Services owns the design, planning, delivery, and operation of all AWS global infrastructure. In other words, we’re the people who keep the cloud running. We support all AWS data centers and all of the servers, storage, networking, power, and cooling equipment that ensure our customers have continual access to the innovation they rely on. We work on the most challenging problems, with thousands of variables impacting the supply chain — and we’re looking for talented people who want to help.

You’ll join a diverse team of software, hardware, and network engineers, supply chain specialists, security experts, operations managers, and other vital roles. You’ll collaborate with people across AWS to help us deliver the highest standards for safety and security while providing seemingly infinite capacity at the lowest possible cost for our customers. And you’ll experience an inclusive culture that welcomes bold ideas and empowers you to own them to completion.


Key job responsibilities
This is an individual contributor role within AWS software development team that owns programmable logic components required for operating and supporting servers in data centers and beyond. The candidate will be an experienced developer who owns end to end programmable logic development that includes design, work estimates, implementation, verification and validation of new features and capabilities. The developer also owns programmable logic component design deployment, operational monitoring, debug and support to improve fleet health. Additionally, the developer is responsible for growing and developing junior team members, and contributing to team growth to improve efficiency and quality.

About the team
Amazon Web Services (AWS) is the world’s most comprehensive and broadly adopted cloud platform. We pioneered cloud computing and never stopped innovating — that’s why customers from the most successful startups to Global 500 companies trust our robust suite of products and services to power their businesses.

Amazon Web Services (AWS) values diverse experiences. Even if you do not meet all of the preferred qualifications and skills listed in the job description, we encourage candidates to apply. If your career is just starting, hasn’t followed a traditional path, or includes alternative experiences, don’t let it stop you from applying.

We value work-life harmony. Achieving success at work should never come at the expense of sacrifices at home, which is why we strive for flexibility as part of our working culture. When we feel supported in the workplace and at home, there’s nothing we can’t achieve in the cloud.

Here at AWS, it’s in our nature to learn and be curious. Our employee-led affinity groups foster a culture of inclusion that empower us to be proud of our differences. Ongoing events and learning experiences, including our Conversations on Race and Ethnicity (CORE) and AmazeCon (gender diversity) conferences, inspire us to never stop embracing our uniqueness.

We’re continuously raising our performance bar as we strive to become Earth’s Best Employer. That’s why you’ll find endless knowledge-sharing, mentorship and other career-advancing resources here to help you develop into a better-rounded professional.

Basic Qualifications
- 3+ years of non-internship professional software development experience
- 3+ years of non-internship design or architecture (design patterns, reliability and scaling) of new and existing systems experience
- Experience programming with at least one software programming language
- Experience with general troubleshooting/debugging of hardware
- 3+ years of programming with at least one hardware description language (Verilog, VHDL) experience

Preferred Qualifications
- 5+ years of non-internship professional software development experience
- 5+ years of full software development life cycle, including coding standards, code reviews, source control management, build processes, testing, and operations experience
- Experience in embedded development in C/C++
- Bachelor's degree or above in electrical engineering, computer engineering, or equivalent
- 5+ years of programming with at least one hardware description language (Verilog, VHDL) experience


"""

from pydantic import BaseModel
class JobD(BaseModel):
    role : str
    requried_skills : list[str]
    preferred_skills : list[str]
    minimum_experience : float |None
    education_requirements : list[str]
    responsibilites : list[str]

jobd_schema = JobD.model_json_schema()


system_prompt = f"""
you are an expert HR assistant 

your job is to analyse job description and extract structured information from them

Return only valid JSON matching this schema:
{jobd_schema}
IMPORTANT:
Do NOT return the schema itself.
Do NOT return fields like "properties", "title" or "type.
Fill the schema with actual infornation extracted fron the job description

If minimum experience is not mentioned, return null.
If infornation for a list is missing, return an enpty list.
Do not invent information.
"""

user_prompt = f"""
Analyze the following job description
{job_description}
"""


message_system = {
    "role": "system",
    "content": system_prompt
}

message_user = {
    "role": "user",
    "content": user_prompt
}

response_format = {
    "type":"json_object"
}

messages=[message_system,message_user]

response = client.chat.completions.create(model= model,messages=messages,response_format=response_format)

answer =response.choices[0].message.content 

raw_json = answer

import json
job_data =json.loads(raw_json)

job = JobD(**job_data)

print(job.minimum_experience)
print(job.education_requirements)



#parse real 
class MatchResult (BaseModel):
    score : float
    details: dict

class Experience (BaseModel):
    company : str | None = None
    role : str | None = None
    duration : str | None = None
    description :str | None = None
    skills_used : list[str] = []


class Resume (BaseModel):     
    name : str | None = None
    email : str | None = None
    phone : str | None = None

    total_experience_years : float | None = None

    skills : list[str] = []
    experience : list[Experience] = []
    education : list[str] = []
    projects : list[str]= []
    certifications : list[str]= []

     
resume_schema = Resume.model_json_schema()

def final_score(job,resume):
    match_schema = MatchResult.model_json_schema()
    prompt = f"""
    You are an HR recruiter.

    Compare the candidate's resume with the job description.

    JOB DESCRIPTION:
    {job.model_dump_json(indent=2)}

    CANDIDATE RESUME:
    {resume.model_dump_json(indent=2)}
    Return JSON matching this schema:

    {match_schema}

    Give me:

    1. Candidate name
    2. Matching skills
    3. Missing important skills
    4. Whether experience requirement is met
    5. Overall match percentage from 0 to 100
    6. A short final verdict

    Keep the response concise and easy to read.
    """
    message={
        "role": "user",
        "content" : prompt
    }
    messages=[message]
    response_format={
        "type": "json_object"
    }
    response = client.chat.completions.create(model=model, messages=messages, response_format=response_format)
    data = json.loads(response.choices[0].message.content)
    return MatchResult(**data)



def parse_resume(resume_text):
    system_prompt = f"""
    You are an expert resume parser.

    Extract information from the resume based on its meaning,
    not only based on exact section headings.

    Different resumes may use different headings.

    For example:
    - Experience
    - Professional Experience
    - Work History
    - Employment
    - Internships

    These may all contain relevant experience.

    Skills may also appear in the skills section, work experience,
    internships or projects.

    Return ONLY valid JSON matching this schema:

    {resume_schema}

    Important rules:

    1. Do not invent information.
    2. If a value is not available, return null.
    3. If a list has no information, return an empty list.
    4. Include internships inside experiences.
    5. Extract skills mentioned across the entire resume.
    """
    user_prompt = f"""
    Parse the following resume:

    {resume_text}
    """
    message_system={
        "role" : "system",
        "content" : system_prompt
    }
    message_user={
        "role" : "user",
        "content" : user_prompt
    }
    messages=[message_system, message_user]
    response_format={
        "type": "json_object"
    }
    response=client.chat.completions.create(model=model, messages=messages, response_format=response_format)
    raw_output = response.choices[0].message.content
    data = json.loads(raw_output)
    resume = Resume(**data)
    return resume




from pypdf import PdfReader
from docx import Document
def read_pdf (file_path):
    reader = PdfReader(file_path)
    text =""
    for page in reader.pages:
        page_text = page.extract_text()
        if page_text :
            text += page_text + "\n"
    return text


def read_docx(file_path):
    document = Document(file_path)
    text = ""
    for paragraph in document.paragraphs:
        if paragraph.text.strip():
            text += paragraph.text + "\n"
    
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    text += cell.text + "\n"
    return text


def read_resume(file_path):
    if file_path.suffix.lower() == ".pdf":
        return read_pdf(file_path)
    elif file_path.suffix.lower() == ".docx":
        return read_docx(file_path)
    else:
        return None



#final part 
from pathlib import Path
resume_folder = Path(__file__).parent / "resumes"
all_results=[]
for file_path in resume_folder.iterdir():
    if file_path.suffix.lower() not in [".pdf", ".docx"]:
        continue
    print("\nProcessing:", file_path.name)
    resume_text = read_resume(file_path)
    parsed_resume=parse_resume(resume_text) # llm call1
    time.sleep(5)
    result = final_score(job, parsed_resume) #llm caLL2
    time.sleep(5)
    print("Score:", result.score)
    all_results.append({
        "name": parsed_resume.name,
        "score": result.score,
        "details": result.details
    })
all_results.sort(
    key=lambda candidate: candidate["score"],
    reverse=True
)
top_2 = all_results[:2]
worst_2 = all_results[-2:]


print("TOP 2 CANDIDATES")
for candidate in top_2:

    print(
        candidate["name"],
        "-",
        candidate["score"],
        "%"
    )

    print(candidate["details"])

print("LOWEST 2 CANDIDATES")
for candidate in worst_2:

    print(
        candidate["name"],
        "-",
        candidate["score"],
        "%"
    )
    print(candidate["details"])
